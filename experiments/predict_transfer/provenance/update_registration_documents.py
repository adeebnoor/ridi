from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import hashlib, json

BASE=Path(__file__).resolve().parent
SOURCE=BASE/'RIDI_PREDICT_v1.1/prereg'
DEST=BASE/'registration_ready/prereg'
name='OSF_Preregistration_RIDI-PREDICT-v1'
olmo='b96024342a77a69aa0dda815c3454a671f477463'
job_text=('After the prespecified eligibility exclusions and SciFact-to-FEVER reallocation, '
          'the study includes NQ (250), HotpotQA (250) and FEVER (300): 800 fresh questions. '
          'This gives 24 model-by-benchmark-by-draw units: 18 predictor units and 6 target units. '
          'No SciFact generation unit is run. The nominal four-benchmark plan would have contained 32 units.')
preflight_text=('Preflight before filing. Tokenizer-only eligibility preflight has completed without '
                'loading model weights or generating answers. The observed counts and exclusion '
                'breakdown are appended in Section 8. The prespecified rules were applied without resampling or changes to thresholds.')
optional_text=('The optional pre-generation AI-judge covariate was not collected. H3 uses the remaining prespecified covariates.')
scope=('Scope. The claim is limited to this exchange, BM25 at k = 10, the three included fresh-panel '
       'benchmarks (NQ, HotpotQA and FEVER), this prompt and these models. The historical motivating '
       'panel and Part A coding sample retain their original four-benchmark scope.')
md=(SOURCE/(name+'.md')).read_text()
oldpre='**Preflight before filing.** Eligibility counts per benchmark are computed at preflight. No model is loaded. The counts are appended to this registration before it is filed.'
changes={oldpre:'**Preflight before filing.** '+preflight_text.split('. ',1)[1],
 'This gives 32 jobs: 24 predictor jobs and 8 target jobs.':job_text,
 'resolved commit, recorded in `LOCK.json`':olmo,
 'If not, H3 uses the remaining covariates, and this is reported.':optional_text,
 '**Scope.** The claim is limited to this exchange, BM25 at k = 10, these four benchmarks, this prompt and these models.':'**Scope.** '+scope.split('. ',1)[1]}
for a,b in changes.items():
    assert md.count(a)==1,a
    md=md.replace(a,b)
insert=(BASE/'preflight_output_final/PREFLIGHT_REGISTRATION_INSERT_EN.md').read_text()
insert=insert.replace('## Tokenizer-only preflight counts (before registration and generation)',
                      '## 8. Completed tokenizer-only preflight before registration')
md+='\n---\n\n'+insert
(DEST/(name+'.md')).write_text(md)

doc=Document(SOURCE/(name+'.docx'))
replacements={
 'Preflight before filing. Eligibility counts per benchmark are computed at preflight. No model is loaded. The counts are appended to this registration before it is filed.':preflight_text,
 'This gives 32 jobs: 24 predictor jobs and 8 target jobs.':job_text,
 'If not, H3 uses the remaining covariates, and this is reported.':optional_text,
 'Scope. The claim is limited to this exchange, BM25 at k = 10, these four benchmarks, this prompt and these models.':scope}
for a,b in replacements.items():
    matches=[p for p in doc.paragraphs if p.text==a];assert len(matches)==1,a
    matches[0].text=b
revision_cells=[c for t in doc.tables for row in t.rows for c in row.cells if 'resolved commit' in c.text]
assert len(revision_cells)==1
revision_cells[0].text=olmo
# Add the actual-count appendix without changing the original scientific sections.
doc.add_page_break()
doc.add_paragraph('8 Completed tokenizer only preflight before registration',
                  style=next(s for s in doc.styles if s.name=='Heading 2'))
paragraphs=insert.split('\n\n')
for p in paragraphs[1:]:
    if p.startswith('| Benchmark'):
        # Narrow source-count breakdown is easier to read as two linked tables in Word.
        count=json.loads((BASE/'preflight_output_final/ELIGIBILITY_COUNTS.json').read_text())['counts']
        for headers,fields in [
            (['Benchmark','Archive','Eligible','Frozen','Known'],['total_queries','registered_eligible','excluded_frozen_panel','excluded_known_source']),
            (['Benchmark','Duplicates','Context','Fresh','Selected'],['excluded_duplicate','excluded_context_length','fresh_eligible','selected_n'])]:
            table=doc.add_table(rows=1, cols=5)
            table.style=next(s for s in doc.styles if s.name=='Table')
            for c,v in zip(table.rows[0].cells,headers):c.text=v
            for key,label in [('nq','NQ'),('hotpotqa','HotpotQA'),('fever','FEVER'),('scifact','SciFact')]:
                row=table.add_row()
                for c,v in zip(row.cells,[label]+[str(count[key][f]) for f in fields]):c.text=v
            for row in table.rows:
                for c in row.cells:
                    for par in c.paragraphs:
                        par.paragraph_format.space_before=Pt(0)
                        par.paragraph_format.space_after=Pt(3)
                        for run in par.runs:run.font.size=Pt(9)
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
    else:
        doc.add_paragraph(p.replace('`',''))
# Long SHA and revision strings in the five-column model table must wrap cleanly.
for row in doc.tables[4].rows:
    for c in row.cells:
        for p in c.paragraphs:
            for run in p.runs:run.font.size=Pt(8)
doc.paragraphs[0].style=next(s for s in doc.styles if s.name=='Title')
for run in doc.paragraphs[0].runs:
    run.font.size=Pt(20)
    run.font.color.rgb=RGBColor(0,0,0)
    run.font.underline=False
for border in list(doc._element.xpath('//w:pBdr')):
    border.getparent().remove(border)
for table in doc.tables:
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
doc.save(DEST/(name+'.docx'))
receipt={'source_md_sha256':hashlib.sha256((SOURCE/(name+'.md')).read_bytes()).hexdigest(),
         'changes':['observed eligibility appendix','OLMo pinned revision','24 actual generation units',
                    'three-benchmark confirmatory scope','optional AI-judge covariate not collected'],
         'hypotheses_thresholds_and_analysis_unchanged':True,'registration_filed':False}
(BASE/'registration_ready/provenance/REGISTRATION_TEXT_UPDATE.json').write_text(json.dumps(receipt,indent=2)+'\n')
print('MD and Word updated; not filed on OSF.')
